from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from app.models.student import Student
from flask import render_template, send_file
from flask_login import login_required
from app.models.transaction import Transaction

from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph
)

from app.reports import reports
from app.models.book import Book


@reports.route("/")
@login_required
def index():
    return render_template("reports/index.html")


@reports.route("/books/pdf")
@login_required
def books_pdf():

    

    books = Book.query.order_by(Book.title).all()

    buffer = BytesIO()

    pdf = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Library Books Report</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph("<br/>", styles["Normal"])
    )

    data = [
        [
            "ID",
            "Title",
            "Author",
            "Category",
            "Available"
        ]
    ]

    for book in books:

        category = (
            book.category.name
            if book.category
            else "-"
        )

        data.append(
            [
                str(book.id),
                book.title,
                book.author,
                category,
                str(book.available)
            ]
        )

    table = Table(data)

    table.setStyle(
        TableStyle(
            [

                ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),

                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

                ("GRID", (0, 0), (-1, -1), 1, colors.black),

                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),

                ("BACKGROUND", (0, 1), (-1, -1), colors.beige),

                ("ALIGN", (0, 0), (-1, -1), "CENTER"),

                ("BOTTOMPADDING", (0, 0), (-1, 0), 10),

            ]
        )
    )

    elements.append(table)

    pdf.build(elements)

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="Library_Books_Report.pdf",
        mimetype="application/pdf"
    )
@reports.route("/students/pdf")
@login_required
def students_pdf():

    students = Student.query.order_by(Student.name).all()

    buffer = BytesIO()

    pdf = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Library Students Report</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph("<br/>", styles["Normal"])
    )

    data = [[
        "ID",
        "Name",
        "Email",
        "Phone"
    ]]

    for student in students:

        data.append([
            str(student.id),
            student.name,
            student.email,
            student.phone
        ])

    table = Table(data)

    table.setStyle(
        TableStyle([

            ("BACKGROUND", (0, 0), (-1, 0), colors.darkgreen),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 10),

        ])
    )

    elements.append(table)

    pdf.build(elements)

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="Library_Students_Report.pdf",
        mimetype="application/pdf"
    )
@reports.route("/transactions/pdf")
@login_required
def transactions_pdf():

    transactions = Transaction.query.order_by(Transaction.id.desc()).all()

    buffer = BytesIO()

    pdf = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Library Transactions Report</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph("<br/>", styles["Normal"])
    )

    data = [[
        "ID",
        "Student",
        "Book",
        "Issue Date",
        "Due Date",
        "Status",
        "Fine"
    ]]

    for t in transactions:

        student = t.student.name if t.student else "-"
        book = t.book.title if t.book else "-"

        issue_date = (
            t.issue_date.strftime("%d-%m-%Y")
            if t.issue_date else "-"
        )

        due_date = (
            t.due_date.strftime("%d-%m-%Y")
            if t.due_date else "-"
        )

        data.append([
            str(t.id),
            student,
            book,
            issue_date,
            due_date,
            t.status,
            f"₹{t.fine:.2f}"
        ])

    table = Table(data)

    table.setStyle(
        TableStyle([

            ("BACKGROUND", (0, 0), (-1, 0), colors.darkred),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 10),

        ])
    )

    elements.append(table)

    pdf.build(elements)

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="Library_Transactions_Report.pdf",
        mimetype="application/pdf"
    )

@reports.route("/books/excel")
@login_required
def books_excel():

    books = Book.query.order_by(Book.title).all()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Books"

    headers = [
        "ID",
        "Title",
        "Author",
        "Category",
        "Available"
    ]

    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(row=1, column=column)

        cell.value = header

        cell.font = Font(bold=True)

    row = 2

    for book in books:

        category = (
            book.category.name
            if book.category
            else "-"
        )

        sheet.cell(row=row, column=1).value = book.id
        sheet.cell(row=row, column=2).value = book.title
        sheet.cell(row=row, column=3).value = book.author
        sheet.cell(row=row, column=4).value = category
        sheet.cell(row=row, column=5).value = book.available

        row += 1

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(

        output,

        as_attachment=True,

        download_name="Library_Books.xlsx",

        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )
@reports.route("/students/excel")
@login_required
def students_excel():

    students = Student.query.order_by(Student.name).all()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Students"

    headers = [
        "ID",
        "Name",
        "Email",
        "Phone"
    ]

    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(row=1, column=column)

        cell.value = header

        cell.font = Font(bold=True)

    row = 2

    for student in students:

        sheet.cell(row=row, column=1).value = student.id
        sheet.cell(row=row, column=2).value = student.name
        sheet.cell(row=row, column=3).value = student.email
        sheet.cell(row=row, column=4).value = student.phone

        row += 1

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Library_Students.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
@reports.route("/transactions/excel")
@login_required
def transactions_excel():

    transactions = Transaction.query.order_by(Transaction.id.desc()).all()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Transactions"

    headers = [
        "ID",
        "Student",
        "Book",
        "Issue Date",
        "Due Date",
        "Status",
        "Fine"
    ]

    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(row=1, column=column)
        cell.value = header
        cell.font = Font(bold=True)

    row = 2

    for transaction in transactions:

        student = (
            transaction.student.name
            if transaction.student
            else "-"
        )

        book = (
            transaction.book.title
            if transaction.book
            else "-"
        )

        issue_date = (
            transaction.issue_date.strftime("%d-%m-%Y")
            if transaction.issue_date
            else "-"
        )

        due_date = (
            transaction.due_date.strftime("%d-%m-%Y")
            if transaction.due_date
            else "-"
        )

        sheet.cell(row=row, column=1).value = transaction.id
        sheet.cell(row=row, column=2).value = student
        sheet.cell(row=row, column=3).value = book
        sheet.cell(row=row, column=4).value = issue_date
        sheet.cell(row=row, column=5).value = due_date
        sheet.cell(row=row, column=6).value = transaction.status
        sheet.cell(row=row, column=7).value = transaction.fine

        row += 1

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Library_Transactions.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )